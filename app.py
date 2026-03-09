from flask import Flask, request, send_file, render_template_string
import re, os, tempfile
from datetime import datetime
import pdfplumber
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

PRICING_TABLE = {
    49.00:74.00,98.00:148.00,47.00:62.00,60.00:78.00,
    100.00:125.00,103.00:130.00,150.00:190.00,193.00:240.00,
    205.00:255.00,283.00:350.00,373.00:460.00,450.00:550.00
}
MARKUP_FALLBACK=1.20
SHIPPING_CHARGE=10.00
C_NAVY="1F4E79"
C_BLUE="2E75B6"
C_WHITE="FFFFFF"
C_ALT="EBF3FB"
C_SUB="D6E4F0"
C_YELLOW="FFF2CC"

HTML = """
<!DOCTYPE html>
<html>
<head>
    <title>Alexander Medical - Pharmacy Parser</title>
    <style>
        * { box-sizing: border-box; margin: 0; padding: 0; }
        body { font-family: -apple-system, BlinkMacSystemFont, sans-serif; background: #f0f4f8; min-height: 100vh; display: flex; align-items: center; justify-content: center; }
        .card { background: white; border-radius: 16px; padding: 48px; max-width: 480px; width: 90%; box-shadow: 0 4px 24px rgba(0,0,0,0.08); }
        .logo { color: #1F4E79; font-size: 22px; font-weight: 700; margin-bottom: 8px; }
        .subtitle { color: #666; font-size: 14px; margin-bottom: 32px; }
        .upload-area { border: 2px dashed #2E75B6; border-radius: 12px; padding: 40px 20px; text-align: center; cursor: pointer; transition: all 0.2s; margin-bottom: 24px; }
        .upload-area:hover { background: #EBF3FB; }
        .upload-area input { display: none; }
        .upload-icon { font-size: 40px; margin-bottom: 12px; }
        .upload-text { color: #2E75B6; font-weight: 600; font-size: 16px; }
        .upload-hint { color: #999; font-size: 13px; margin-top: 4px; }
        .filename { color: #1F4E79; font-size: 14px; margin-top: 8px; font-weight: 500; }
        .btn { width: 100%; padding: 14px; background: #1F4E79; color: white; border: none; border-radius: 8px; font-size: 16px; font-weight: 600; cursor: pointer; transition: background 0.2s; }
        .btn:hover { background: #2E75B6; }
        .btn:disabled { background: #ccc; cursor: not-allowed; }
        .status { margin-top: 16px; padding: 12px; border-radius: 8px; font-size: 14px; text-align: center; display: none; }
        .status.processing { background: #EBF3FB; color: #1F4E79; display: block; }
        .status.error { background: #ffeaea; color: #c00; display: block; }
        .footer { margin-top: 24px; text-align: center; color: #bbb; font-size: 12px; }
    </style>
</head>
<body>
<div class="card">
    <div class="logo">Alexander Medical</div>
    <div class="subtitle">Val Vista Pharmacy Report Parser</div>
    <form method="POST" action="/upload" enctype="multipart/form-data" id="form">
        <div class="upload-area" onclick="document.getElementById('file').click()">
            <div class="upload-icon">📄</div>
            <div class="upload-text">Click to select PDF</div>
            <div class="upload-hint">Val Vista Pharmacy A/R Report</div>
            <div class="filename" id="fname"></div>
            <input type="file" id="file" name="file" accept=".pdf" onchange="document.getElementById('fname').textContent=this.files[0].name">
        </div>
        <button class="btn" type="submit" id="btn">Generate Excel Report</button>
        <div class="status processing" id="status" style="display:none">⏳ Processing PDF, please wait...</div>
    </form>
    {% if error %}
    <div class="status error">{{ error }}</div>
    {% endif %}
    <div class="footer">Alexander Medical © 2026</div>
</div>
<script>
document.getElementById('form').onsubmit = function() {
    document.getElementById('btn').disabled = true;
    document.getElementById('status').style.display = 'block';
}
</script>
</body>
</html>
"""

def fix_amount(s):
    s=str(s).replace(",","").replace("$","").strip()
    if re.match(r"^\d+$",s):
        v=int(s)
        if v>1000:
            return float(str(v)[:-2]+"."+str(v)[-2:])
        return float(v)
    try:
        return float(s)
    except:
        return 0.0

def lookup_price(cost):
    if cost in PRICING_TABLE:
        return float(PRICING_TABLE[cost])
    return round(cost*MARKUP_FALLBACK,2)

def cl(n):
    return get_column_letter(n)

def make_header(ws,row,n,bg=C_NAVY,fc=C_WHITE):
    for c in range(1,n+1):
        x=ws.cell(row,c)
        x.fill=PatternFill("solid",fgColor=bg)
        x.font=Font(color=fc,bold=True,size=10)
        x.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)

def make_border(ws,r1,r2,c1,c2):
    t=Side(style="thin",color="CCCCCC")
    for row in ws.iter_rows(r1,r2,c1,c2):
        for c in row:
            c.border=Border(top=t,left=t,right=t,bottom=t)

def make_title(ws,txt,ncols):
    ws.merge_cells("A1:"+cl(ncols)+"1")
    ws["A1"]=txt
    ws["A1"].font=Font(bold=True,size=13,color=C_WHITE)
    ws["A1"].fill=PatternFill("solid",fgColor=C_NAVY)
    ws["A1"].alignment=Alignment(horizontal="center",vertical="center")
    ws.row_dimensions[1].height=32
    ws.sheet_view.showGridLines=False

def parse_pdf(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        full_text="\n".join(p.extract_text() or "" for p in pdf.pages)
    if not full_text.strip():
        try:
            import pytesseract
            from pdf2image import convert_from_path
            from PIL import Image
            pages=convert_from_path(pdf_path,dpi=300)
            full_text="\n".join(pytesseract.image_to_string(p.convert("L"),config="--psm 6") for p in pages)
        except:
            pass
    meta={
        "report_start":find_re(r"Start Date:\s*(\S+)",full_text),
        "report_end":find_re(r"End Date:\s*(\S+)",full_text),
        "print_date":find_re(r"Print Date:\s*(\S+)",full_text),
    }
    txns=[]
    shipping=[]
    current_pt=None
    lines=full_text.split("\n")
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:
            i+=1
            continue
        if re.search(r"payment posted",line,re.I):
            i+=1
            continue
        if re.search(r"out of state",line,re.I):
            if not re.search(r"return",line,re.I):
                dm=re.match(r"(\d{1,2}/\d{1,2}/\d{4})",line)
                if dm:
                    am=re.search(r"\$(\d+\.\d{2})",line)
                    shipping.append({"date":dm.group(1),"description":line,"amount":float(am.group(1)) if am else SHIPPING_CHARGE})
            i+=1
            continue
        pt_m=re.match(r"^([A-Z][a-z'\-]+(?:\s+[A-Z][a-z'\-]+)?,\s+[A-Z][a-zA-Z'\-].*?)(?:\s+[/\\].*)?$",line)
        if pt_m and not re.match(r"^\d{1,2}/\d{1,2}/\d{4}",line) and "Patient Subtotal" not in line and "Balance Forward" not in line and "Date Description" not in line:
            cand=pt_m.group(1).strip()
            if re.match(r"^[A-Z][a-z].*,\s+[A-Z]",cand):
                current_pt=cand
                i+=1
                continue
        tx=re.match(r"^(\d{1,2}/\d{1,2}/\d{4})\s+(.+?)\s+(\d+)\s+\$?(\d[\d,]*\.?\d*)\s+\$?(\d[\d,]*\.?\d*)\s+\$?(\d[\d,]*\.?\d*)\s+[\$\|]?\s*(\d[\d,]*\.?\d*)",line)
        if tx:
            date_s,desc,qty,deb,crd,tax,tot=tx.groups()
            deb_f=fix_amount(deb)
            crd_f=fix_amount(crd)
            if deb_f==0.0 and crd_f>0.0:
                i+=1
                continue
            if deb_f<1.0:
                i+=1
                continue
            combined=line+" "+(lines[i+1] if i+1<len(lines) else "")
            df_m=re.search(r"DateFilled:\s*([A-Za-z]+ \d+ \d{4})",combined)
            rx_m=re.match(r"^(\d{6}-\d+)",desc)
            med=re.sub(r"^\d{6}-\d+\s*[-—]\s*","",desc).strip()
            med=re.sub(r"\s*[-—]?\s*[Qq]tv?:\d.*$","",med,flags=re.I).strip()
            clinic=lookup_price(deb_f)
            txns.append({
                "report_date":date_s,
                "patient_name":current_pt or "UNKNOWN",
                "rx_number":rx_m.group(1) if rx_m else "",
                "medication":med,
                "qty":int(qty),
                "date_filled":df_m.group(1) if df_m else "",
                "pharmacy_cost":deb_f,
                "credit":crd_f,
                "tax":fix_amount(tax),
                "net_pharm_cost":deb_f-crd_f,
                "clinic_price":clinic,
                "gross_profit":round(clinic-deb_f,2),
            })
        i+=1
    return txns,shipping,meta

def find_re(pattern,text,default=""):
    m=re.search(pattern,text)
    return m.group(1) if m else default

def build_workbook(txns,shipping,meta,out_path):
    df=pd.DataFrame(txns) if txns else pd.DataFrame()
    dfs=pd.DataFrame(shipping) if shipping else pd.DataFrame()
    wb=Workbook()
    wb.remove(wb.active)
    build_raw_sheet(wb,df,meta)
    build_weekly_sheet(wb,df,dfs,meta)
    build_patient_sheet(wb,df,dfs)
    build_med_sheet(wb,df)
    wb.save(out_path)

def build_raw_sheet(wb,df,meta):
    ws=wb.create_sheet("Raw Data")
    make_title(ws,"ALEXANDER MEDICAL  |  Raw Transaction Data",14)
    ws.freeze_panes="A4"
    ws.merge_cells("A2:N2")
    ws["A2"]="Period: "+meta.get("report_start","")+" to "+meta.get("report_end","")+"  |  Print Date: "+meta.get("print_date","")
    ws["A2"].font=Font(italic=True,size=9,color="555555")
    ws["A2"].alignment=Alignment(horizontal="center")
    hdrs=["Report Date","Patient Name","Rx #","Medication","Qty","Date Filled","Pharm Cost","Credit","Tax","Net Pharm Cost","Clinic Price","Gross Profit","Margin %","Notes"]
    for c,h in enumerate(hdrs,1):
        ws.cell(3,c,h)
    make_header(ws,3,len(hdrs))
    ws.row_dimensions[3].height=34
    if not df.empty:
        keys=["report_date","patient_name","rx_number","medication","qty","date_filled","pharmacy_cost","credit","tax","net_pharm_cost","clinic_price","gross_profit"]
        for i,row in df.iterrows():
            r=i+4
            alt=PatternFill("solid",fgColor=C_ALT if i%2==0 else C_WHITE)
            for c,k in enumerate(keys,1):
                x=ws.cell(r,c,row.get(k,""))
                x.fill=alt
                x.font=Font(size=9)
                x.alignment=Alignment(vertical="center")
            mg=(row["gross_profit"]/row["clinic_price"]) if row.get("clinic_price") else 0
            ws.cell(r,13,mg).number_format="0.0%"
            ws.cell(r,13).fill=alt
            ws.cell(r,14,"").fill=alt
            for c in [7,8,9,10,11,12]:
                ws.cell(r,c).number_format='"$"#,##0.00'
        make_border(ws,3,len(df)+3,1,14)
    widths=[11,24,13,42,5,13,12,9,7,14,12,12,9,18]
    for c,w in enumerate(widths,1):
        ws.column_dimensions[cl(c)].width=w

def build_weekly_sheet(wb,df,dfs,meta):
    ws=wb.create_sheet("Weekly Summary")
    make_title(ws,"WEEKLY SUMMARY DASHBOARD",7)
    ws.merge_cells("A2:G2")
    ws["A2"]="Period: "+meta.get("report_start","")+" to "+meta.get("report_end","")+"  |  Generated: "+datetime.today().strftime("%m/%d/%Y")
    ws["A2"].font=Font(italic=True,size=9,color="555555")
    ws["A2"].alignment=Alignment(horizontal="center")
    for c in range(1,8):
        ws.column_dimensions[cl(c)].width=17
    if df.empty:
        ws["A4"]="No data parsed."
        return
    tot_rx=len(df)
    tot_cost=df["pharmacy_cost"].sum()
    tot_rev=df["clinic_price"].sum()
    tot_prof=df["gross_profit"].sum()
    tot_marg=(tot_prof/tot_rev*100) if tot_rev else 0
    tot_ship=dfs["amount"].sum() if not dfs.empty else 0
    oos_ct=len(dfs) if not dfs.empty else 0
    kpi_labels=["Total Rx","Pharm Cost","Clinic Revenue","Gross Profit","Margin %","OOS Patients","Shipping Rev"]
    kpi_values=[tot_rx,tot_cost,tot_rev,tot_prof,tot_marg,oos_ct,tot_ship]
    kpi_fmts=["#","$","$","$","%","#","$"]
    ws.row_dimensions[4].height=18
    ws.row_dimensions[5].height=38
    for c,(lbl,val,fmt) in enumerate(zip(kpi_labels,kpi_values,kpi_fmts),1):
        lc=ws.cell(4,c,lbl)
        lc.fill=PatternFill("solid",fgColor=C_BLUE)
        lc.font=Font(bold=True,size=9,color=C_WHITE)
        lc.alignment=Alignment(horizontal="center",vertical="center")
        vc=ws.cell(5,c,val)
        vc.font=Font(bold=True,size=16,color=C_BLUE)
        vc.alignment=Alignment(horizontal="center",vertical="center")
        if fmt=="$":
            vc.number_format='"$"#,##0.00'
        elif fmt=="%":
            vc.number_format='0.0"%"'
        else:
            vc.number_format="#,##0"
    day_hdrs=["Date","Rx Count","Pharm Cost","Clinic Revenue","Gross Profit","OOS Shipping","Total Revenue"]
    r0=7
    ws.row_dimensions[r0].height=24
    for c,h in enumerate(day_hdrs,1):
        ws.cell(r0,c,h)
    make_header(ws,r0,len(day_hdrs),bg=C_BLUE)
    daily=df.groupby("report_date").agg(rx_count=("pharmacy_cost","count"),pharm=("pharmacy_cost","sum"),rev=("clinic_price","sum"),prof=("gross_profit","sum")).reset_index().sort_values("report_date")
    sbd={}
    if not dfs.empty:
        for _,s in dfs.iterrows():
            sbd[s["date"]]=sbd.get(s["date"],0)+s["amount"]
    r=r0+1
    for _,day in daily.iterrows():
        alt=PatternFill("solid",fgColor=C_ALT if r%2==0 else C_WHITE)
        sh=sbd.get(day["report_date"],0)
        vals=[day["report_date"],day["rx_count"],day["pharm"],day["rev"],day["prof"],sh,day["rev"]+sh]
        for c,v in enumerate(vals,1):
            cell=ws.cell(r,c,v)
            cell.fill=alt
            cell.font=Font(size=10)
            cell.alignment=Alignment(horizontal="center" if c<=2 else "right",vertical="center")
        for c in [3,4,5,6,7]:
            ws.cell(r,c).number_format='"$"#,##0.00'
        r+=1
    ws.cell(r,1,"TOTAL").font=Font(bold=True)
    for c,v in zip([2,3,4,5,6,7],[daily["rx_count"].sum(),daily["pharm"].sum(),daily["rev"].sum(),daily["prof"].sum(),sum(sbd.values()),daily["rev"].sum()+sum(sbd.values())]):
        cell=ws.cell(r,c,v)
        cell.font=Font(bold=True)
        cell.fill=PatternFill("solid",fgColor=C_SUB)
    for c in [3,4,5,6,7]:
        ws.cell(r,c).number_format='"$"#,##0.00'
    make_border(ws,r0,r,1,7)
    sr=r+3
    ws.merge_cells("A"+str(sr)+":G"+str(sr))
    ws["A"+str(sr)]="OUT-OF-STATE SHIPPING DETAIL"
    ws["A"+str(sr)].font=Font(bold=True,size=11,color=C_WHITE)
    ws["A"+str(sr)].fill=PatternFill("solid",fgColor=C_BLUE)
    ws["A"+str(sr)].alignment=Alignment(horizontal="center",vertical="center")
    ws.row_dimensions[sr].height=24
    sr+=1
    for c,h in enumerate(["Date","Description","Amount"],1):
        ws.cell(sr,c,h)
    make_header(ws,sr,3,bg=C_BLUE)
    r2=sr+1
    if not dfs.empty:
        for _,s in dfs.iterrows():
            ws.cell(r2,1,s["date"])
            ws.cell(r2,2,s["description"])
            ws.cell(r2,3,s["amount"])
            ws.cell(r2,3).number_format='"$"#,##0.00'
            r2+=1
        ws.cell(r2,2,"TOTAL").font=Font(bold=True)
        ws.cell(r2,3,dfs["amount"].sum())
        ws.cell(r2,3).number_format='"$"#,##0.00'
        ws.cell(r2,3).font=Font(bold=True)
        make_border(ws,sr,r2,1,3)
    else:
        ws.cell(r2,1,"No out-of-state shipping charges this period.")

def build_patient_sheet(wb,df,dfs):
    ws=wb.create_sheet("Patient Summary")
    make_title(ws,"PATIENT SUMMARY  |  Yellow = Out-of-State",8)
    ws.freeze_panes="A3"
    hdrs=["Patient Name","# Rx","Pharm Cost","Clinic Revenue","Gross Profit","Margin %","OOS Shipping","Total Revenue"]
    for c,h in enumerate(hdrs,1):
        ws.cell(2,c,h)
    make_header(ws,2,8)
    ws.row_dimensions[2].height=32
    if df.empty:
        return
    pt_ship={}
    if not dfs.empty:
        for _,s in dfs.iterrows():
            desc=s["description"].upper()
            for pn in df["patient_name"].unique():
                if not pn or pn=="UNKNOWN":
                    continue
                parts=pn.split(",")
                last=parts[0].strip().upper()
                first=parts[1].strip().split()[0].upper() if len(parts)>1 and parts[1].strip() else ""
                init=(first[0]+last[0]) if first and last else ""
                if (last in desc) or (init and init in desc):
                    pt_ship[pn]=pt_ship.get(pn,0)+s["amount"]
                    break
    summary=df.groupby("patient_name").agg(rx=("pharmacy_cost","count"),cost=("pharmacy_cost","sum"),rev=("clinic_price","sum"),prof=("gross_profit","sum")).reset_index().sort_values("patient_name")
    r=3
    for _,pt in summary.iterrows():
        alt=PatternFill("solid",fgColor=C_ALT if r%2==1 else C_WHITE)
        ship=pt_ship.get(pt["patient_name"],0)
        totrev=pt["rev"]+ship
        marg=(pt["prof"]/pt["rev"]*100) if pt["rev"] else 0
        vals=[pt["patient_name"],pt["rx"],pt["cost"],pt["rev"],pt["prof"],marg,ship,totrev]
        for c,v in enumerate(vals,1):
            cell=ws.cell(r,c,v)
            cell.fill=alt
            cell.font=Font(size=10)
            cell.alignment=Alignment(vertical="center")
        for c in [3,4,5,7,8]:
            ws.cell(r,c).number_format='"$"#,##0.00'
        ws.cell(r,6).number_format='0.0"%"'
        if ship>0:
            for c in range(1,9):
                ws.cell(r,c).fill=PatternFill("solid",fgColor=C_YELLOW)
        r+=1
    ws.cell(r,1,"GRAND TOTAL").font=Font(bold=True)
    total_ship=sum(pt_ship.values())
    for c,v in zip([2,3,4,5,7,8],[summary["rx"].sum(),summary["cost"].sum(),summary["rev"].sum(),summary["prof"].sum(),total_ship,summary["rev"].sum()+total_ship]):
        cell=ws.cell(r,c,v)
        cell.font=Font(bold=True)
        cell.fill=PatternFill("solid",fgColor=C_SUB)
    total_marg=(summary["prof"].sum()/summary["rev"].sum()*100) if summary["rev"].sum() else 0
    ws.cell(r,6,total_marg)
    ws.cell(r,6).number_format='0.0"%"'
    ws.cell(r,6).font=Font(bold=True)
    for c in [3,4,5,7,8]:
        ws.cell(r,c).number_format='"$"#,##0.00'
    make_border(ws,2,r,1,8)
    widths=[26,7,14,16,14,10,14,15]
    for c,w in enumerate(widths,1):
        ws.column_dimensions[cl(c)].width=w

def build_med_sheet(wb,df):
    ws=wb.create_sheet("Medication Summary")
    make_title(ws,"MEDICATION SUMMARY",7)
    hdrs=["Medication","Times Dispensed","Total Pharm Cost","Total Clinic Revenue","Total Gross Profit","Avg Margin %","Avg Clinic Price"]
    for c,h in enumerate(hdrs,1):
        ws.cell(2,c,h)
    make_header(ws,2,7)
    ws.row_dimensions[2].height=34
    if df.empty:
        return
    df2=df.copy()
    df2["med"]=df2["medication"].str.replace(r"^\d{6}-\d+\s*[-—]\s*","",regex=True)
    df2["med"]=df2["med"].str.replace(r"\s*[-—]?\s*[Qq]tv?:\d.*$","",regex=True,flags=re.IGNORECASE).str.strip()
    meds=df2.groupby("med").agg(count=("pharmacy_cost","count"),cost=("pharmacy_cost","sum"),rev=("clinic_price","sum"),prof=("gross_profit","sum"),avg_price=("clinic_price","mean")).reset_index().sort_values("count",ascending=False)
    r=3
    for _,m in meds.iterrows():
        alt=PatternFill("solid",fgColor=C_ALT if r%2==1 else C_WHITE)
        marg=(m["prof"]/m["rev"]*100) if m["rev"] else 0
        vals=[m["med"],m["count"],m["cost"],m["rev"],m["prof"],marg,m["avg_price"]]
        for c,v in enumerate(vals,1):
            cell=ws.cell(r,c,v)
            cell.fill=alt
            cell.font=Font(size=10)
            cell.alignment=Alignment(vertical="center")
        for c in [3,4,5,7]:
            ws.cell(r,c).number_format='"$"#,##0.00'
        ws.cell(r,6).number_format='0.0"%"'
        r+=1
    make_border(ws,2,r-1,1,7)
    widths=[52,16,18,20,18,13,16]
    for c,w in enumerate(widths,1):
        ws.column_dimensions[cl(c)].width=w

@app.route("/", methods=["GET"])
def index():
    return render_template_string(HTML, error=None)

@app.route("/upload", methods=["POST"])
def upload():
    if "file" not in request.files:
        return render_template_string(HTML, error="No file selected.")
    f=request.files["file"]
    if not f.filename.endswith(".pdf"):
        return render_template_string(HTML, error="Please upload a PDF file.")
    with tempfile.TemporaryDirectory() as tmpdir:
        pdf_path=os.path.join(tmpdir,"input.pdf")
        out_path=os.path.join(tmpdir,"output.xlsx")
        f.save(pdf_path)
        txns,ship,meta=parse_pdf(pdf_path)
        build_workbook(txns,ship,meta,out_path)
        filename=f.filename.replace(".pdf","_processed.xlsx")
        return send_file(out_path,as_attachment=True,download_name=filename,mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

if __name__=="__main__":
    port=int(os.environ.get("PORT",5000))
    app.run(host="0.0.0.0",port=port)
